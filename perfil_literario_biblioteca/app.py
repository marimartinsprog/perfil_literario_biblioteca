from __future__ import annotations

import csv
import io
import json
import os
import re
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Iterable

import requests
from flask import Flask, jsonify, redirect, render_template, request, session, url_for
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
ACERVO_PATH = DATA_DIR / "acervo.json"
META_PATH = DATA_DIR / "meta.json"

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "troque-esta-chave-em-producao")
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024

ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "biblioteca2026")
ALLOWED_EXTENSIONS = {"csv", "xlsx"}

TAG_KEYWORDS = {
    "fantasia": ["fantasia", "magia", "bruxa", "bruxo", "dragao", "castelo", "elfo", "fada", "feitiço", "feitico", "mitologia", "reino"],
    "ficcao_cientifica": ["ficcao cientifica", "robo", "espaco", "futuro", "tecnologia", "alien", "distopia", "galaxia", "android", "maquina do tempo"],
    "misterio": ["misterio", "crime", "investigacao", "detetive", "suspense", "policial", "assassinato", "enigma", "desaparecida", "segredo", "caso"],
    "romance": ["romance", "amor", "paixao", "casal", "namoro", "coracao", "noiva", "noivo", "beijo", "felizes para nunca", "altar"],
    "aventura": ["aventura", "jornada", "viagem", "exploracao", "heroi", "ilha", "labirinto", "batalha", "mar", "tesouro", "expedicao"],
    "terror": ["terror", "horror", "medo", "fantasma", "sombrio", "morte", "zumbi", "cemiterio", "assombrada", "monstro", "pesadelo"],
    "biografia": ["biografia", "autobiografia", "historia de vida", "memorias", "diario de", "minha vida"],
    "historia": ["historia", "passado", "guerra", "imperio", "seculo", "revolucao", "brasil", "nazismo", "auschwitz", "medieval"],
    "filosofia": ["filosofia", "etica", "pensamento", "existencialismo", "platao", "aristoteles", "nietzsche", "sabedoria", "sentido da vida"],
    "psicologia": ["psicologia", "mente", "comportamento", "terapia", "freud", "jung", "emocao", "ansiedade", "trauma", "cerebro"],
    "autoajuda": ["autoajuda", "desenvolvimento pessoal", "sucesso", "motivacao", "inspirador", "habitos", "lideranca", "felicidade", "proposito"],
    "juvenil": ["juvenil", "jovem", "adolescente", "escola", "colegio", "garota", "garoto", "princesa", "diario de um banana"],
    "infantil": ["infantil", "crianca", "fabula", "contos", "bonequinha", "turma", "bichos", "brincadeira", "porquinho", "porquinhos", "fabulas", "ninar"],
    "conhecimento": ["ciencia", "cientista", "matematica", "fisica", "universo", "descoberta", "enciclopedia", "manual", "guia", "documentario"],
    "classico": ["classico", "machado de assis", "shakespeare", "dostoievski", "tolstoi", "eca de queiroz", "jane austen", "victor hugo", "oscar wilde"],
    "poesia": ["poesia", "poema", "versos", "rima", "poetica", "sonetos"],
}

AUTHOR_TAGS = {
    "christie, agatha": {"misterio", "classico"},
    "mcfadden, freida": {"misterio", "suspense"},
    "donlea, charlie": {"misterio", "suspense"},
    "king, stephen": {"terror", "misterio"},
    "poe, edgar allan": {"terror", "misterio", "classico"},
    "lovecraft": {"terror", "classico"},
    "hoover, colleen": {"romance", "sentimental"},
    "quinn, julia": {"romance", "historia"},
    "moyes, jojo": {"romance", "sentimental"},
    "sparks, nicholas": {"romance", "sentimental"},
    "hazelwood, ali": {"romance"},
    "han, jenny": {"romance", "juvenil"},
    "reid, taylor jenkins": {"romance"},
    "meyer, stephenie": {"fantasia", "romance", "juvenil"},
    "riordan, rick": {"fantasia", "aventura", "juvenil"},
    "lewis, c. s.": {"fantasia", "aventura", "classico"},
    "tolkien": {"fantasia", "aventura", "classico"},
    "rowling": {"fantasia", "aventura", "juvenil"},
    "garber, stephanie": {"fantasia", "romance", "juvenil"},
    "cass, kiera": {"romance", "juvenil", "fantasia"},
    "verne, julio": {"aventura", "ficcao_cientifica", "classico"},
    "orwell, george": {"ficcao_cientifica", "classico", "filosofia"},
    "asimov": {"ficcao_cientifica", "classico"},
    "bradbury": {"ficcao_cientifica", "classico"},
    "clarke, arthur": {"ficcao_cientifica", "classico"},
    "assis, machado de": {"classico", "romance", "filosofia"},
    "queiroz, eca de": {"classico", "romance", "historia"},
    "dostoievski": {"classico", "filosofia", "psicologia"},
    "tolstoi": {"classico", "historia", "filosofia"},
    "austen, jane": {"classico", "romance"},
    "amado, jorge": {"classico", "historia", "romance"},
    "coelho, paulo": {"filosofia", "autoajuda"},
    "alves, rubem": {"filosofia", "autoajuda", "poesia"},
    "cortella": {"filosofia", "autoajuda"},
    "karnal": {"historia", "filosofia"},
    "quintana, mario": {"poesia", "classico"},
    "furnari, eva": {"infantil", "juvenil", "fantasia"},
    "kinney, jeff": {"juvenil", "infantil", "aventura"},
    "montgomery, l. m": {"juvenil", "romance", "classico"},
    "montes, raphael": {"misterio", "terror"},
}

ADVANCED_AUTHORS = [
    "assis, machado", "dostoievski", "tolstoi", "nietzsche", "platao", "aristoteles",
    "orwell, george", "queiroz, eca", "shakespeare", "woolf", "kafka", "joyce",
    "hugo, victor", "camus", "sartre", "foucault", "bauman", "freud", "jung",
]

BEGINNER_HINTS = {
    "infantil", "juvenil", "poesia", "autoajuda", "misterio", "romance"
}
ADVANCED_HINTS = {"classico", "filosofia", "historia", "psicologia"}

PROFILE_CONFIG = {
    "Aventureiro": {
        "slug": "aventureiro",
        "emoji": "🧭",
        "description": "Você lê em busca de movimento, descobertas e grandes jornadas.",
        "tags": {"aventura", "jornada", "viagem", "exploracao", "heroi"},
    },
    "Sonhador": {
        "slug": "sonhador",
        "emoji": "✨",
        "description": "Sua imaginação abre portais para mundos mágicos e extraordinários.",
        "tags": {"fantasia", "magia", "bruxa", "dragao", "mitologia"},
    },
    "Analítico": {
        "slug": "analitico",
        "emoji": "🔬",
        "description": "Você gosta de ideias, lógica, ciência e futuros possíveis.",
        "tags": {"ficcao_cientifica", "tecnologia", "futuro", "psicologia"},
    },
    "Sentimental": {
        "slug": "sentimental",
        "emoji": "🧡",
        "description": "Histórias ganham força quando despertam emoções e conexões humanas.",
        "tags": {"romance", "amor", "sentimental", "psicologia"},
    },
    "Misterioso": {
        "slug": "misterioso",
        "emoji": "🔍",
        "description": "Você segue pistas, suspeita de tudo e precisa descobrir a verdade.",
        "tags": {"misterio", "crime", "investigacao", "suspense", "terror"},
    },
    "Reflexivo": {
        "slug": "reflexivo",
        "emoji": "🧘",
        "description": "Você procura leituras que provoquem reflexão e transformem o olhar.",
        "tags": {"filosofia", "psicologia", "autoajuda", "poesia", "classico"},
    },
    "Curioso": {
        "slug": "curioso",
        "emoji": "📚",
        "description": "Você lê para compreender pessoas, épocas, fatos e o mundo real.",
        "tags": {"biografia", "historia", "conhecimento", "classico"},
    },
}


def normalize(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", value.lower()).strip()


def contains_keyword(text: str, keyword: str) -> bool:
    """Evita falsos positivos como 'amor' dentro de 'metamorfose'."""
    normalized_keyword = normalize(keyword)
    pattern = rf"(?<![a-z0-9]){re.escape(normalized_keyword)}(?![a-z0-9])"
    return re.search(pattern, text) is not None


def generate_book_tags(title: str, author: str) -> set[str]:
    text = normalize(f"{title} {author}")
    tags: set[str] = set()

    for tag, keywords in TAG_KEYWORDS.items():
        if any(contains_keyword(text, keyword) for keyword in keywords):
            tags.add(tag)

    author_norm = normalize(author)
    for author_key, author_tags in AUTHOR_TAGS.items():
        if normalize(author_key) in author_norm:
            tags.update(author_tags)

    # Inferências leves por termos comuns, mantendo os livros sem classificação
    # com uma tag neutra e score menor.
    if not tags:
        if any(contains_keyword(text, word) for word in ["manual", "guia", "introducao", "curso", "aprenda", "tecnicas"]):
            tags.update({"conhecimento", "autoajuda"})
        elif any(contains_keyword(text, word) for word in ["recontadas", "ninar"]):
            tags.add("infantil")
        elif any(contains_keyword(text, word) for word in ["cronicas", "crônicas"]):
            tags.update({"classico", "poesia"})
        elif any(contains_keyword(text, word) for word in ["contos"]):
            tags.add("poesia")
        else:
            tags.add("geral")

    return tags


def classify_level(title: str, author: str, tags: Iterable[str]) -> str:
    text = normalize(f"{title} {author}")
    tags = set(tags)

    explicit_child_hints = [
        "infantil", "crianca", "ninar", "porquinho", "porquinhos",
        "bonequinha", "turma", "diario de um banana", "recontadas",
    ]
    if any(contains_keyword(text, hint) for hint in explicit_child_hints):
        return "iniciante"
    if any(normalize(author_name) in text for author_name in ADVANCED_AUTHORS):
        return "avancado"
    if tags & {"infantil", "juvenil"}:
        return "iniciante"
    if len(title) > 82 and tags & ADVANCED_HINTS:
        return "avancado"
    if len(tags & ADVANCED_HINTS) >= 2 and not tags & {"autoajuda", "poesia"}:
        return "avancado"
    if tags & BEGINNER_HINTS and len(title) <= 62:
        return "iniciante"
    return "intermediario"


def build_book(title: str, author: str, code: str | None = None) -> dict:
    title = str(title or "").strip()
    author = str(author or "Autor não informado").strip() or "Autor não informado"
    tags = generate_book_tags(title, author)
    return {
        "titulo": title,
        "autor": author,
        "codigo": str(code or "").strip(),
        "tags": sorted(tags),
        "nivel": classify_level(title, author, tags),
    }


def load_collection() -> list[dict]:
    if not ACERVO_PATH.exists():
        return []
    with ACERVO_PATH.open("r", encoding="utf-8") as file:
        return json.load(file)


def save_collection(books: list[dict], source: str) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with ACERVO_PATH.open("w", encoding="utf-8") as file:
        json.dump(books, file, ensure_ascii=False, indent=2)
    meta = {
        "total": len(books),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "source": source,
    }
    with META_PATH.open("w", encoding="utf-8") as file:
        json.dump(meta, file, ensure_ascii=False, indent=2)


def level_compatibility(reader_level: str, book_level: str) -> float:
    exact_bonus = {"iniciante": 0.45, "intermediario": 0.35, "avancado": 0.55}
    if reader_level == book_level:
        return exact_bonus[reader_level]
    matrix = {
        ("iniciante", "intermediario"): 0.02,
        ("iniciante", "avancado"): -0.35,
        ("intermediario", "iniciante"): 0.08,
        ("intermediario", "avancado"): 0.10,
        ("avancado", "intermediario"): 0.12,
        ("avancado", "iniciante"): -0.15,
    }
    return matrix.get((reader_level, book_level), 0.0)


def jaccard(a: set[str], b: set[str]) -> float:
    union = a | b
    return len(a & b) / len(union) if union else 0.0


def canonical_author(author: str) -> str:
    """Agrupa variações como 'ROWLING, J. K.' e 'ROWLING, Joanne K'."""
    normalized = normalize(author)
    if "," in normalized:
        return normalized.split(",", 1)[0].strip()
    parts = normalized.split()
    return parts[-1] if parts else normalized


def choose_profile(profile_scores: dict[str, int], user_tags: set[str]) -> str:
    if profile_scores:
        highest = max(profile_scores.values())
        tied = [name for name, score in profile_scores.items() if score == highest]
    else:
        tied = list(PROFILE_CONFIG)

    if len(tied) == 1:
        return tied[0]

    return max(
        tied,
        key=lambda name: len(PROFILE_CONFIG[name]["tags"] & user_tags),
    )


def recommend_books(level: str, user_tags: set[str], profile: str, limit: int = 5) -> list[dict]:
    books = load_collection()
    expanded_tags = set(user_tags) | PROFILE_CONFIG[profile]["tags"]
    scored: list[dict] = []

    for book in books:
        book_tags = set(book.get("tags", []))
        overlap = expanded_tags & book_tags
        score = jaccard(expanded_tags, book_tags)
        score += level_compatibility(level, book.get("nivel", "intermediario"))
        score += min(len(overlap), 3) * 0.08
        if not overlap:
            score -= 0.50
        if book_tags == {"conhecimento"} and "conhecimento" not in expanded_tags:
            score -= 0.12

        scored.append({
            **book,
            "score": round(score, 4),
            "afinidades": sorted(overlap),
        })

    scored.sort(key=lambda item: (item["score"], len(item["afinidades"]), item["titulo"]), reverse=True)

    selected: list[dict] = []
    author_counts: Counter[str] = Counter()

    # Primeira passagem: autores totalmente diferentes.
    for book in scored:
        author_key = canonical_author(book["autor"])
        if author_counts[author_key] == 0:
            selected.append(book)
            author_counts[author_key] += 1
        if len(selected) == limit:
            break

    # Segunda passagem: em nichos limitados, aceita no máximo dois do mesmo autor.
    if len(selected) < limit:
        selected_titles = {book["titulo"] for book in selected}
        for book in scored:
            author_key = canonical_author(book["autor"])
            if book["titulo"] in selected_titles or author_counts[author_key] >= 2:
                continue
            selected.append(book)
            author_counts[author_key] += 1
            if len(selected) == limit:
                break

    return selected


def read_uploaded_collection(file_storage) -> list[dict]:
    filename = secure_filename(file_storage.filename or "")
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Envie um arquivo CSV ou XLSX.")

    content = file_storage.read()
    raw_rows: list[tuple[str, str, str]] = []

    if extension == "csv":
        text = content.decode("utf-8-sig", errors="replace")
        sample = text[:4096]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        for row in reader:
            normalized = {normalize(key): value for key, value in row.items() if key}
            title = normalized.get("titulo") or normalized.get("título") or ""
            author = normalized.get("autor") or ""
            code = normalized.get("cod") or normalized.get("codigo") or normalized.get("cód") or ""
            if str(title).strip():
                raw_rows.append((str(title), str(author), str(code)))
    else:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [normalize(value) for value in next(rows)]
        indexes = {header: index for index, header in enumerate(headers)}
        title_index = indexes.get("titulo")
        author_index = indexes.get("autor")
        code_index = indexes.get("cod", indexes.get("codigo"))
        if title_index is None or author_index is None:
            raise ValueError("A planilha precisa ter as colunas TÍTULO e AUTOR.")
        for row in rows:
            title = row[title_index] if title_index < len(row) else ""
            author = row[author_index] if author_index < len(row) else ""
            code = row[code_index] if code_index is not None and code_index < len(row) else ""
            if str(title or "").strip():
                raw_rows.append((str(title), str(author or ""), str(code or "")))

    if not raw_rows:
        raise ValueError("Nenhum livro válido foi encontrado.")

    return [build_book(title, author, code) for title, author, code in raw_rows]


def google_sheet_csv_url(link: str) -> str:
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", link)
    if not match:
        raise ValueError("Link do Google Sheets inválido.")
    sheet_id = match.group(1)
    gid_match = re.search(r"gid=([0-9]+)", link)
    gid = gid_match.group(1) if gid_match else "0"
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"


@app.get("/")
def index():
    return render_template("index.html", profiles=PROFILE_CONFIG)


@app.post("/api/recomendar")
def recommend_api():
    payload = request.get_json(silent=True) or {}
    level = payload.get("nivel")
    user_tags = set(payload.get("tags") or [])
    profile_scores = payload.get("perfis") or {}

    if level not in {"iniciante", "intermediario", "avancado"}:
        return jsonify({"erro": "Nível de leitura inválido."}), 400
    if not user_tags:
        return jsonify({"erro": "Responda ao questionário antes de gerar o resultado."}), 400

    profile = choose_profile(profile_scores, user_tags)
    recommendations = recommend_books(level, user_tags, profile)
    config = PROFILE_CONFIG[profile]

    return jsonify({
        "perfil": profile,
        "slug": config["slug"],
        "emoji": config["emoji"],
        "descricao": config["description"],
        "nivel": level,
        "tags": sorted(user_tags),
        "livros": recommendations,
    })


@app.get("/api/status")
def status_api():
    books = load_collection()
    meta = {}
    if META_PATH.exists():
        meta = json.loads(META_PATH.read_text(encoding="utf-8"))
    return jsonify({
        "total": len(books),
        "updated_at": meta.get("updated_at"),
        "source": meta.get("source", "Acervo inicial"),
    })


@app.route("/admin", methods=["GET", "POST"])
def admin():
    error = None
    success = session.pop("admin_success", None)

    if request.method == "POST" and request.form.get("action") == "login":
        if request.form.get("password") == ADMIN_PASSWORD:
            session["admin"] = True
            return redirect(url_for("admin"))
        error = "Senha incorreta."

    meta = {}
    if META_PATH.exists():
        meta = json.loads(META_PATH.read_text(encoding="utf-8"))

    return render_template(
        "admin.html",
        authenticated=session.get("admin", False),
        error=error,
        success=success,
        total=len(load_collection()),
        meta=meta,
    )


@app.post("/admin/upload")
def admin_upload():
    if not session.get("admin"):
        return redirect(url_for("admin"))
    try:
        upload = request.files.get("arquivo")
        if not upload or not upload.filename:
            raise ValueError("Selecione um arquivo.")
        books = read_uploaded_collection(upload)
        save_collection(books, f"Upload: {secure_filename(upload.filename)}")
        session["admin_success"] = f"Acervo atualizado com {len(books)} livros."
    except Exception as exc:  # mensagem amigável no painel
        session["admin_success"] = f"Erro: {exc}"
    return redirect(url_for("admin"))


@app.post("/admin/sheets")
def admin_sheets():
    if not session.get("admin"):
        return redirect(url_for("admin"))
    try:
        link = request.form.get("sheets_url", "").strip()
        export_url = google_sheet_csv_url(link)
        response = requests.get(export_url, timeout=20)
        response.raise_for_status()
        fake_file = type("Uploaded", (), {
            "filename": "google_sheets.csv",
            "read": lambda self: response.content,
        })()
        books = read_uploaded_collection(fake_file)
        save_collection(books, "Google Sheets")
        session["admin_success"] = f"Sincronização concluída: {len(books)} livros."
    except Exception as exc:
        session["admin_success"] = f"Erro: {exc}"
    return redirect(url_for("admin"))


@app.post("/admin/logout")
def admin_logout():
    session.clear()
    return redirect(url_for("admin"))


if __name__ == "__main__":
    app.run(debug=True)
